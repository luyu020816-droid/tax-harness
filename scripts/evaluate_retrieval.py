"""Frozen development/holdout evaluation using real local model inference.

Separate databases, immutable output directories, no production data resets.
The cache stores actual profile-bound inference; timing labels disclose reuse.
"""

import argparse
import hashlib
import io
import json
import math
import re
import subprocess
import time
import zipfile
from pathlib import Path

import psycopg
from app.cli import initialize
from app.config import ROOT, Settings
from app.db import Database
from app.evidence_selection import (
    SELECTION_VERSION,
    SUPPORT_SELECTION_VERSION,
    load_policy,
    support_score,
)
from app.identity import Identity
from app.knowledge import Knowledge
from app.knowledge_jobs import KnowledgeJobs
from app.knowledge_retrieval import KnowledgeRetrieval
from app.retrieval_engine import LocalRetrieval
from app.retrieval_metrics import calibrate, ndcg, paired_family_interval, recall_units
from app.values import digest, safe_json
from docx import Document
from psycopg import sql
from psycopg.conninfo import make_conninfo
from retrieval_boundaries import run_boundaries
from retrieval_model_checks import run_model_checks

AUTH = Identity("U001")


def write(path, value):
    path.write_text(json.dumps(safe_json(value), ensure_ascii=False, indent=2), encoding="utf-8")


class CachedInference:
    """Exact-input cache of real model outputs. Never supplies synthetic vectors."""

    def __init__(self, engine):
        self.engine, self.vectors, self.scores = engine, {}, {}
        self.counts = {
            "real_embedding_texts": 0,
            "cached_embedding_texts": 0,
            "real_rerank_pairs": 0,
            "cached_rerank_pairs": 0,
        }

    def profile(self):
        return self.engine.profile()

    def external_work(self):
        return self.engine.external_work()

    def count(self, text):
        return self.engine.count(text)

    def embed(self, texts):
        missing = list(dict.fromkeys(t for t in texts if t not in self.vectors))
        self.counts["cached_embedding_texts"] += len(texts) - len(missing)
        if missing:
            self.counts["real_embedding_texts"] += len(missing)
            values = self.engine.embed(missing)
            self.vectors.update(zip(missing, values, strict=True))
        return [self.vectors[t] for t in texts]

    def rerank(self, query, passages):
        key = digest([query, passages])
        if isinstance(self.scores.get(key), list):
            self.counts["cached_rerank_pairs"] += len(passages)
            self.scores.update(
                (digest([query, p]), score)
                for p, score in zip(passages, self.scores[key], strict=True)
            )
            return self.scores[key]
        missing = list(dict.fromkeys(p for p in passages if digest([query, p]) not in self.scores))
        self.counts["cached_rerank_pairs"] += len(passages) - len(missing)
        if missing:
            self.counts["real_rerank_pairs"] += len(missing)
            values = self.engine.rerank(query, missing)
            self.scores.update(
                (digest([query, p]), score) for p, score in zip(missing, values, strict=True)
            )
        return [self.scores[digest([query, p])] for p in passages]


def document_bytes(item):
    if item["format"] == "md":
        return item["text"].encode()
    doc = Document()
    doc.add_heading(item["title"], 0)
    doc.add_heading("第1条 办理要求", 1)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "证据单元"
    table.cell(0, 1).text = "完整要求（金额单位：元；比例按百分比）"
    for suffix in ("A", "C", "B"):
        cells = table.add_row().cells
        cells[0].text = "[" + item["id"] + suffix + "]"
        cells[1].text = item["units"][item["id"] + suffix]
    doc.add_heading("第2条 联系与版本", 1)
    doc.add_paragraph("[" + item["id"] + "D] " + item["units"][item["id"] + "D"])
    result = io.BytesIO()
    doc.save(result)
    normalized = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(result.getvalue())) as source,
        zipfile.ZipFile(normalized, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for name in sorted(source.namelist()):
            entry = zipfile.ZipInfo(name, date_time=(2020, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(entry, source.read(name))
    return normalized.getvalue()


def database(settings, split, output):
    name = (
        "tax_harness_retrieval_"
        + ("dev" if split == "development" else "holdout")
        + "_v1_"
        + digest(str(output.resolve()))[:8]
    )
    with psycopg.connect(
        make_conninfo(settings.admin_database_url, dbname="postgres"), autocommit=True
    ) as conn:
        if not conn.execute("SELECT 1 FROM pg_database WHERE datname=%s", (name,)).fetchone():
            conn.execute(sql.SQL("CREATE DATABASE {}").format(sql.Identifier(name)))
    settings = settings.model_copy(
        update={
            "database_url": make_conninfo(settings.database_url, dbname=name),
            "admin_database_url": make_conninfo(settings.admin_database_url, dbname=name),
        }
    )
    initialize(settings)
    db = Database(settings.database_url)
    db.open()
    return settings, db


def ingest(store, jobs, items, dataset_sha, settings):
    result = {}
    for item in items:
        created = store.upload(
            AUTH,
            name=item["id"] + "." + item["format"],
            content=document_bytes(item),
            company_id="C001",
            title=item["title"],
            request_id="retrieval-v1:" + dataset_sha + ":" + item["id"],
        )
        if store.get(AUTH, created["document_id"])["active_version_id"] != created["version_id"]:
            claimed = jobs.claim()
            assert claimed and str(claimed["job_id"]) == created["job_id"]
            jobs.run(claimed)
        current = store.get(AUTH, created["document_id"])
        assert current["active_version_id"] == created["version_id"], store.job(
            AUTH, created["job_id"]
        )
        result[item["id"]] = created
        # Isolated corpus fixtures represent documents uploaded over many days.
        # Never disable or change the production upload rate limiter.
        with psycopg.connect(settings.admin_database_url) as conn:
            conn.execute(
                "UPDATE knowledge_files SET created_at='2026-09-01'::timestamptz WHERE file_id=%s",
                (created["file_id"],),
            )
        print("indexed", item["id"], flush=True)
    return result


def known_chunks(db, documents):
    with db.transaction(readonly=True) as conn:
        rows = conn.execute(
            "SELECT c.chunk_id,c.version_id,c.content,c.locator,c.block_ordinals,d.title,v.document_id FROM knowledge_chunks c JOIN knowledge_versions v USING(version_id) JOIN knowledge_documents d USING(document_id) WHERE v.version_id=d.active_version_id AND d.status='active'"
        ).fetchall()
        block_units = {
            (str(r["version_id"]), r["ordinal"]): re.findall(r"\[([DH]\d{2}[ABCD])\]", r["content"])
            for r in conn.execute(
                "SELECT version_id,ordinal,content FROM knowledge_blocks"
            ).fetchall()
        }
    by_document = {str(v["document_id"]): k for k, v in documents.items()}
    return {
        str(r["chunk_id"]): {
            **safe_json(r),
            "fixture_document": by_document.get(str(r["document_id"])),
            "units": list(
                set(
                    u
                    for ordinal in r["block_ordinals"]
                    for u in block_units.get((str(r["version_id"]), ordinal), [])
                )
            ),
        }
        for r in rows
    }


def query_trace(db, query_id):
    with db.transaction(readonly=True) as conn:
        return safe_json(
            conn.execute(
                "SELECT candidates,coverage,timings,retrieval_mode FROM retrieval_queries WHERE query_id=%s",
                (query_id,),
            ).fetchone()
        )


def measure(question, bundle, trace, chunks, items):
    document = question.get("document")
    required = question["required_units"]
    units = {key: r["units"] for key, r in chunks.items()}
    grades = (
        {
            key: max([items[document]["grades"].get(u, 0) for u in r["units"]], default=0)
            if r["fixture_document"] == document
            else 0
            for key, r in chunks.items()
        }
        if document
        else {}
    )
    ranks = {
        **trace["coverage"].get("candidate_channels", {}),
        "reranked": [c["chunk_id"] for c in trace["candidates"]],
    }
    row = {
        "id": question["id"],
        "family": question["family"],
        "answerable": question["answerable"],
        "status": bundle["status"],
        "degraded_reason": bundle["degraded_reason"],
        "mode": bundle["retrieval_mode"],
        "query_id": bundle["query_id"],
    }
    for name, ids in ranks.items():
        row[name + "_recall20"] = recall_units(ids, required, units)
        row[name + "_ndcg10"] = ndcg(ids, grades) if document else None
    served = "\n".join(e["excerpt"] for e in bundle["evidence"])
    row["complete_evidence"] = (
        all(
            "[" + unit + "]" in served and items[document]["units"][unit] in served
            for unit in required
        )
        if required
        else None
    )
    row["selected_evidence"] = len(bundle["evidence"])
    row["evidence_tokens"] = bundle["coverage"].get("evidence_tokens", 0)
    row["budget_incomplete"] = bundle["coverage"].get("budget_incomplete", False)
    legacy_ids, seen, remaining = [], set(), 3000
    qualified = []
    for candidate in trace["candidates"]:
        c = chunks[candidate["chunk_id"]]
        cost = math.ceil(len((c["content"] + c["title"] + str(c["locator"])).encode()) / 2)
        sha = text_sha256(c["content"])
        if sha not in seen and cost <= remaining and len(legacy_ids) < 6:
            legacy_ids.append(candidate["chunk_id"])
            seen.add(sha)
            remaining -= cost
        if (
            bundle["coverage"].get("selection_cutoff") is not None
            and support_score(
                candidate,
                {
                    "version": bundle["coverage"].get("selection_version"),
                    "similarity_weight": bundle["coverage"].get("selection_similarity_weight", 0),
                    "rerank_floor": bundle["coverage"].get("selection_rerank_floor"),
                },
            )
            >= bundle["coverage"]["selection_cutoff"]
        ):
            qualified.append(candidate["chunk_id"])

    def complete(ids):
        text = "\n".join(chunks[key]["content"] for key in ids)
        return (
            all(
                "[" + unit + "]" in text and items[document]["units"][unit] in text
                for unit in required
            )
            if required
            else None
        )

    row["legacy_topk_complete_evidence"] = complete(legacy_ids)
    row["unexpanded_filtered_complete_evidence"] = (
        complete(qualified) if bundle["coverage"].get("selection_cutoff") is not None else None
    )
    row["citation_checks"] = all(
        e["excerpt_hash"] == text_sha256(e["excerpt"]) and e["locator"] for e in bundle["evidence"]
    )
    row["retained_unknown_evidence"] = (
        bool(bundle["evidence"]) if not question["answerable"] else None
    )
    row["strongest_candidate"] = max(
        (c.get("rerank_logit", -math.inf) for c in trace["candidates"]), default=-100.0
    )
    row["strongest_gold_seed"] = (
        max(
            (
                c.get("rerank_logit", -math.inf)
                for c in trace["candidates"]
                if set(units.get(c["chunk_id"], [])) & set(required)
            ),
            default=-100.0,
        )
        if required
        else None
    )
    return row


def text_sha256(text):
    # Excerpt hashes are SHA256 of UTF-8 text, not the JSON digest helper.
    return hashlib.sha256(text.encode()).hexdigest()


def summary(rows):
    positives = [r for r in rows if r["answerable"]]
    negatives = [r for r in rows if not r["answerable"]]
    result = {"questions": len(rows), "answerable": len(positives), "unanswerable": len(negatives)}
    for key in (
        "lexical_recall20",
        "dense_recall20",
        "rrf_recall20",
        "reranked_recall20",
        "lexical_ndcg10",
        "dense_ndcg10",
        "rrf_ndcg10",
        "reranked_ndcg10",
        "complete_evidence",
        "legacy_topk_complete_evidence",
        "unexpanded_filtered_complete_evidence",
        "citation_checks",
    ):
        values = [r[key] for r in positives if r.get(key) is not None]
        result[key] = sum(values) / len(values) if values else None
    result["unanswerable_retained_evidence"] = sum(
        r["retained_unknown_evidence"] for r in negatives
    )
    result["unanswerable_no_evidence_rate"] = (
        sum(not r["retained_unknown_evidence"] for r in negatives) / len(negatives)
        if negatives
        else None
    )
    result["paired_recall_rrf_minus_lexical"] = paired_family_interval(
        positives, "rrf_recall20", "lexical_recall20"
    )
    result["paired_ndcg_reranked_minus_rrf"] = paired_family_interval(
        positives, "reranked_ndcg10", "rrf_ndcg10"
    )
    return result


def main(args):
    args.output.mkdir(parents=True, exist_ok=False)
    code_snapshot = {
        p.name: p.read_text(encoding="utf-8") for p in sorted((ROOT / "backend/app").glob("*.py"))
    }
    evaluator_sha = hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
    write(args.output / "code-snapshot.json", code_snapshot)
    suite = json.loads((args.suite_directory / f"{args.split}.json").read_text(encoding="utf-8"))
    suite_sha = digest(suite)
    write(args.output / "frozen-suite.json", suite)
    settings, db = database(Settings(), args.split, args.output)
    engine = CachedInference(LocalRetrieval(settings.knowledge_model_directory))
    store, jobs = Knowledge(db, settings), KnowledgeJobs(db, settings, engine)
    profile_id, manifest = engine.profile()
    write(args.output / "model-profile.json", {"profile_id": profile_id, "manifest": manifest})
    cache_source = None
    if args.inference_cache:
        raw = args.inference_cache.read_bytes()
        cached = json.loads(raw)
        assert cached["profile_id"] == profile_id, "Inference cache profile mismatch"
        engine.vectors = cached["vectors"]
        engine.scores = cached["rerank_scores"]
        cache_source = {
            "path": str(args.inference_cache),
            "sha256": hashlib.sha256(raw).hexdigest(),
        }
    if args.split == "holdout":
        policy, error = load_policy(args.policy, profile_id)
        assert policy, error
        development = json.loads(
            (args.suite_directory / "development.json").read_text(encoding="utf-8")
        )
        assert policy["development_sha256"] == digest(development), (
            "Development fixture does not match locked policy"
        )
        write(args.output / "locked-policy.json", policy)
    else:
        policy = None
    try:
        documents = ingest(store, jobs, suite["documents"], suite_sha, settings)
        chunks = known_chunks(db, documents)
        write(args.output / "corpus-chunks.json", chunks)
        questions = suite["questions"] + [q for q in suite["boundary_questions"] if "query" in q]
        prime = time.monotonic()
        engine.embed([q["query"] for q in questions])
        priming_ms = round((time.monotonic() - prime) * 1000)
        baseline = KnowledgeRetrieval(db, settings, engine, selection_policy=False)
        items = {d["id"]: d for d in suite["documents"]}
        baseline_rows = []
        for question in questions:
            bundle = baseline.search(AUTH, company_id="C001", query=question["query"])
            row = measure(question, bundle, query_trace(db, bundle["query_id"]), chunks, items)
            baseline_rows.append(row)
            write(
                args.output / (question["id"] + "-baseline.json"),
                {
                    "question": question,
                    "measurement": row,
                    "bundle": bundle,
                    "trace": query_trace(db, bundle["query_id"]),
                },
            )
            print("measured", question["id"], flush=True)
        if args.split == "development":
            cutoff, fit = calibrate(
                [r["strongest_gold_seed"] for r in baseline_rows if r["answerable"]],
                [r["strongest_candidate"] for r in baseline_rows if not r["answerable"]],
            )
            raw = {
                "version": SELECTION_VERSION,
                "profile_id": profile_id,
                "mode": "hybrid",
                "development_sha256": suite_sha,
                "cutoff": cutoff,
                "calibration": fit,
            }
            if args.joint_support:
                trials = []
                for weight in (0, 4, 8, 16, 32, 64):
                    for floor in (-8, -6, -4, 0):
                        signal = {
                            "version": SUPPORT_SELECTION_VERSION,
                            "similarity_weight": weight,
                            "rerank_floor": floor,
                        }
                        positive, negative = [], []
                        for q in questions:
                            record = json.loads(
                                (args.output / (q["id"] + "-baseline.json")).read_text(
                                    encoding="utf-8"
                                )
                            )
                            scored = [
                                support_score(c, signal)
                                for c in record["trace"]["candidates"]
                                if not q["answerable"]
                                or set(chunks[c["chunk_id"]]["units"]) & set(q["required_units"])
                            ]
                            (positive if q["answerable"] else negative).append(
                                max(scored, default=-100.0)
                            )
                        threshold, stats = calibrate(positive, negative)
                        trials.append({**signal, "cutoff": threshold, "calibration": stats})
                best = min(
                    trials,
                    key=lambda t: (
                        t["calibration"]["development_false_accepts"],
                        -t["calibration"]["development_gold_seed_retained"],
                        t["similarity_weight"],
                        -t["rerank_floor"],
                    ),
                )
                raw.update(best)
                write(args.output / "calibration-trials.json", trials)
            write(args.output / "policy.json", raw)
            policy, error = load_policy(args.output / "policy.json", profile_id)
            assert policy, error
        selected = KnowledgeRetrieval(db, settings, engine, selection_policy=policy)
        final_rows = []
        for question in questions:
            bundle = selected.search(AUTH, company_id="C001", query=question["query"])
            for e in bundle["evidence"]:
                assert selected.read(AUTH, e["citation_id"])["excerpt"] == e["excerpt"]
            row = measure(question, bundle, query_trace(db, bundle["query_id"]), chunks, items)
            final_rows.append(row)
            write(
                args.output / (question["id"] + "-selected.json"),
                {
                    "question": question,
                    "measurement": row,
                    "bundle": bundle,
                    "trace": query_trace(db, bundle["query_id"]),
                },
            )
        write(
            args.output / "real-inference-cache.json",
            {"profile_id": profile_id, "vectors": engine.vectors, "rerank_scores": engine.scores},
        )
        model_result = (
            run_model_checks(settings, db, selected, suite, args.output)
            if args.split == "holdout" and args.live_model
            else None
        )
        if model_result:
            write(args.output / "model-checks.json", model_result)
        boundary_result = (
            run_boundaries(settings, store, jobs, selected, documents, suite_sha)
            if args.split == "holdout"
            else None
        )
        if boundary_result:
            write(args.output / "boundaries.json", boundary_result)
        result = {
            "split": args.split,
            "suite_sha256": suite_sha,
            "policy": policy,
            "profile_id": profile_id,
            "baseline": summary(baseline_rows),
            "selected": summary(final_rows),
            "measurements": final_rows,
            "boundaries": boundary_result,
            "model_checks": model_result,
            "inference_cache_source": cache_source,
            "inference_counts": engine.counts,
            "fixture_metadata_note": "Synthetic historical file creation timestamps in isolated evaluation DB; runtime quotas unchanged.",
            "evaluator_sha256": evaluator_sha,
            "query_embedding_priming_ms": priming_ms,
            "timing_note": "Query vectors and exact rerank pairs reuse actual profile-bound outputs; warm replay is not production request latency. Baseline may also hit imported caches; inference_counts separately records actual model work. Query priming is reported separately.",
            "gold_review": "implementer-authored synthetic unit labels; independent human semantic review pending",
            "historical_30_questions": "out_of_scope_for_this_batch",
            "code_commit": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT)
            .decode()
            .strip(),
            "code_fingerprint": digest(code_snapshot),
        }
        write(args.output / "report.json", result)
        print(
            json.dumps(
                {
                    "split": args.split,
                    "selected": result["selected"],
                    "policy_id": policy["policy_id"],
                },
                ensure_ascii=False,
            ),
            flush=True,
        )
    finally:
        engine.engine.close()
        db.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--split", required=True, choices=["development", "holdout"])
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--policy", type=Path)
    parser.add_argument("--inference-cache", type=Path)
    parser.add_argument("--live-model", action="store_true")
    parser.add_argument("--suite-directory", type=Path, default=ROOT / "evals/retrieval-v1")
    parser.add_argument("--joint-support", action="store_true")
    args = parser.parse_args()
    if args.split == "holdout" and not args.policy:
        parser.error("Holdout requires the locked development policy")
    main(args)
