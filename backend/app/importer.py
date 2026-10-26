import hashlib
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from app.errors import AppError
from app.values import jsonb, uid


def money(value) -> Decimal:
    try:
        result = Decimal(str(value))
        if not result.is_finite() or result < 0:
            raise ValueError
        return result.quantize(Decimal("0.01"))
    except (ValueError, InvalidOperation) as exc:
        raise AppError("IMPORT_INVALID", "明细金额无效。") from exc


def import_samples(conn, workbook: Path, rules_file: Path) -> dict:
    source_hash = hashlib.sha256(workbook.read_bytes() + rules_file.read_bytes()).hexdigest()
    existing = conn.execute(
        "SELECT * FROM import_batches WHERE source_hash=%s", (source_hash,)
    ).fetchone()
    if existing:
        return {**existing, "already_imported": True}
    if conn.execute("SELECT 1 FROM invoices LIMIT 1").fetchone():
        raise AppError(
            "IMPORT_CONFLICT", "当前数据库已有业务资料，禁止覆盖。请使用独立演示/测试库。", 409
        )
    w = load_workbook(workbook, data_only=True, read_only=True)
    counts = {"companies": 0, "users": 0, "invoices": 0, "bank_entries": 0, "rules": 0}
    for row in w["企业与用户"].values:
        if isinstance(row[0], str) and re.fullmatch(r"C\d+", row[0]):
            conn.execute(
                "INSERT INTO companies VALUES(%s,%s,%s)",
                (row[0], row[1], jsonb([row[0], row[1], row[2]])),
            )
            conn.execute("INSERT INTO company_heads(company_id) VALUES(%s)", (row[0],))
            counts["companies"] += 1
        if isinstance(row[0], str) and re.fullmatch(r"U\d+", row[0]):
            conn.execute("INSERT INTO users(user_id,display_name) VALUES(%s,%s)", (row[0], row[1]))
            for company, permission in zip(["C001", "C002"], row[2:4], strict=True):
                conn.execute(
                    "INSERT INTO user_company_grants(user_id,company_id,can_read,can_update) VALUES(%s,%s,%s,%s)",
                    (
                        row[0],
                        company,
                        permission in ("可查看并修改", "仅查看"),
                        permission == "可查看并修改",
                    ),
                )
            counts["users"] += 1
    for line, row in enumerate(w["发票"].values, 1):
        if not isinstance(row[0], str) or not re.fullmatch(r"C\d+", row[0]):
            continue
        company, number, day, seller, item, amount, purpose, note = row[:8]
        if not isinstance(day, datetime | date):
            raise AppError("IMPORT_INVALID", f"发票第 {line} 行日期无效。")
        conn.execute(
            "INSERT INTO invoices(company_id,invoice_id,invoice_date,seller,item,amount,purpose,note) VALUES(%s,%s,%s,%s,%s,%s,%s,%s)",
            (
                company,
                number,
                day.date() if isinstance(day, datetime) else day,
                seller,
                item,
                money(amount),
                purpose,
                note,
            ),
        )
        counts["invoices"] += 1
    for line, row in enumerate(w["银行流水"].values, 1):
        if not isinstance(row[0], str) or not re.fullmatch(r"C\d+", row[0]):
            continue
        company, number, day, direction, amount, counterparty, note = row[:7]
        if (
            not isinstance(day, datetime | date)
            or direction not in ("收入", "支出")
            or money(amount) <= 0
        ):
            raise AppError("IMPORT_INVALID", f"流水第 {line} 行字段无效。")
        conn.execute(
            "INSERT INTO bank_entries VALUES(%s,%s,%s,%s,%s,%s,%s)",
            (
                company,
                number,
                day.date() if isinstance(day, datetime) else day,
                "income" if direction == "收入" else "expense",
                money(amount),
                counterparty,
                note,
            ),
        )
        counts["bank_entries"] += 1
    document = Document(rules_file)
    sections = {}
    key = None
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        match = re.match(r"(R\d{3}).*?(\d+\.\d+)", text)
        if match:
            key = match.groups()
            sections[key] = []
        if key and text:
            sections[key].append(text)
    for cells in document.tables[0].rows[1:]:
        code, name, version, scope, effective, status = [c.text.strip() for c in cells.cells]
        condition = {"kind": "money_summary"} if code == "R001" else {"kind": "bank_read_only"}
        topic = "amounts" if code == "R001" else "bank"
        if code in ("R002", "R003"):
            topic = "purpose"
            condition = (
                {"kind": "purpose_all"}
                if (code, version) == ("R002", "2.0")
                else {"kind": "purpose_threshold", "operator": "gte", "threshold": "1000.00"}
            )
        rule_id = uid()
        conn.execute(
            "INSERT INTO rule_versions(rule_version_id,rule_code,display_name,business_version,effective_from,status,topic,full_text,condition_json,source_ref) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",
            (
                rule_id,
                code,
                name,
                version,
                date.fromisoformat(effective),
                "disabled" if "停用" in status else "active",
                topic,
                "\n".join(sections[(code, version)]),
                jsonb(condition),
                f"03_业务规则.docx#{code}-{version}",
            ),
        )
        companies = ["C001"] if code == "R002" else ["C002"] if code == "R003" else ["C001", "C002"]
        for company in companies:
            conn.execute("INSERT INTO rule_scopes VALUES(%s,%s)", (rule_id, company))
        counts["rules"] += 1
    w.close()
    if counts != {"companies": 2, "users": 2, "invoices": 40, "bank_entries": 48, "rules": 5}:
        raise AppError("IMPORT_INVALID", "材料数量不符合 v1.1.0，整个批次已拒绝。", details=counts)
    batch = {
        "batch_id": uid(),
        "source_hash": source_hash,
        "baseline_id": "competition-v1.1.0",
        "row_counts": counts,
    }
    conn.execute(
        "INSERT INTO import_batches(batch_id,source_hash,baseline_id,row_counts) VALUES(%s,%s,%s,%s)",
        (batch["batch_id"], source_hash, batch["baseline_id"], jsonb(counts)),
    )
    return batch
