"""Bounded XLSX parsing. This module never accesses the model or database."""

import hashlib
import re
import zipfile
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from xml.etree.ElementTree import ParseError, iterparse

from openpyxl import load_workbook

from app.errors import AppError
from app.values import INVOICE_ID_PATTERN

MAX_FILE_BYTES = 5 * 1024 * 1024
PARSER_VERSION = "1"
ALIASES = {
    "企业编号": "company_id",
    "发票编号": "invoice_id",
    "发票号码": "invoice_id",
    "日期": "invoice_date",
    "发票日期": "invoice_date",
    "销售方": "seller",
    "销售方名称": "seller",
    "项目名称": "item",
    "项目": "item",
    "金额": "amount",
    "价税合计": "amount",
    "用途": "purpose",
    "发票用途": "purpose",
    "备注": "note",
}
REQUIRED = {"invoice_id", "invoice_date", "seller", "item", "amount"}


def validate_archive(content: bytes):
    if not content or len(content) > MAX_FILE_BYTES:
        raise AppError("UPLOAD_TOO_LARGE", "文件不能为空，且最大为 5 MiB。", 413)
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            if (
                len(entries) > 2000
                or sum(entry.file_size for entry in entries) > 50 * 1024 * 1024
                or "[Content_Types].xml" not in names
                or "xl/workbook.xml" not in names
                or any(entry.flag_bits & 1 for entry in entries)
                or any(
                    any(
                        part in name.lower()
                        for part in ("vbaproject", "externallinks/", "embeddings/", "macrosheets/")
                    )
                    for name in names
                )
            ):
                raise ValueError
            content_types = archive.read("[Content_Types].xml").lower()
            if b"macroenabled" in content_types or b"vbaproject" in content_types:
                raise ValueError
            for name in names:
                if (
                    name.endswith(".rels")
                    or name.startswith("xl/worksheets/")
                    and name.endswith(".xml")
                ):
                    with archive.open(name) as stream:
                        for _, element in iterparse(stream, events=("end",)):
                            tag = element.tag.rsplit("}", 1)[-1]
                            if (
                                tag == "f"
                                or element.attrib.get("TargetMode") == "External"
                                or tag == "row"
                                and int(element.attrib.get("r", "0")) > 10000
                            ):
                                raise ValueError
                            if tag == "c":
                                column = re.match("[A-Z]+", element.attrib.get("r", ""))
                                if column and (len(column[0]) > 1 or column[0] > "P"):
                                    raise ValueError
                            element.clear()
    except (ValueError, ParseError, zipfile.BadZipFile, RuntimeError) as exc:
        raise AppError(
            "WORKBOOK_UNSUPPORTED", "请使用未加密、不含宏和外部内容的标准 XLSX 模板。", 422
        ) from exc


def _text(value, limit, *, required=False):
    if value is None and not required:
        return None
    if not isinstance(value, str) or len(value) > limit or (required and not value.strip()):
        raise ValueError(f"文本格式无效或超过 {limit} 字符")
    return value.strip() if required else value


def _normalize(values, company_id):
    if values.get("company_id") != company_id:
        raise AppError(
            "IMPORT_VALIDATION_FAILED", "文件企业编号必须与当前选择的企业一致，整批未处理。", 422
        )
    number = _text(values.get("invoice_id"), 32, required=True)
    if not re.fullmatch(INVOICE_ID_PATTERN, number):
        raise ValueError("发票编号须为文本，格式为 F 加 3—12 位数字，不能丢失前导零")
    day = values.get("invoice_date")
    if isinstance(day, datetime):
        if day.time().isoformat() != "00:00:00":
            raise ValueError("日期不能包含非零时间")
        day = day.date()
    elif isinstance(day, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", day):
        day = date.fromisoformat(day)
    if not isinstance(day, date):
        raise ValueError("日期须为 Excel 日期或 YYYY-MM-DD")
    amount = values.get("amount")
    if isinstance(amount, bool):
        raise ValueError("金额格式无效")
    try:
        amount = Decimal(str(amount))
        if (
            not amount.is_finite()
            or amount < 0
            or amount.as_tuple().exponent < -2
            or amount >= Decimal("10000000000000000")
        ):
            raise ValueError
    except (ValueError, InvalidOperation) as exc:
        raise ValueError("金额须为非负、最多两位小数的十进制数") from exc
    return {
        "company_id": company_id,
        "invoice_id": number,
        "invoice_date": day.isoformat(),
        "seller": _text(values.get("seller"), 200, required=True),
        "item": _text(values.get("item"), 200, required=True),
        "amount": format(amount, ".2f"),
        "purpose": _text(values.get("purpose"), 2000),
        "note": _text(values.get("note"), 2000),
    }


def parse_invoice_workbook(content: bytes, company_id: str) -> list[dict]:
    validate_archive(content)
    workbook = None
    try:
        workbook = load_workbook(
            BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        if "发票" not in workbook.sheetnames:
            raise AppError("WORKBOOK_UNSUPPORTED", "缺少“发票”工作表，请下载模板。", 422)
        for sheet in workbook:
            if sheet.title not in ("发票", "说明"):
                for cells in sheet.iter_rows(max_row=10001):
                    if any(cell.value is not None for cell in cells):
                        raise AppError(
                            "WORKBOOK_UNSUPPORTED",
                            "只支持发票数据和说明页，不接受其他业务工作表。",
                            422,
                        )
        sheet = workbook["发票"]
        if sheet.max_column and sheet.max_column > 16:
            raise AppError("WORKBOOK_UNSUPPORTED", "工作表列数超过模板范围。", 422)
        headers = next(sheet.iter_rows(max_row=1), ())
        mapping = []
        for cell in headers:
            if cell.value is None:
                mapping.append(None)
                continue
            if (
                cell.data_type == "f"
                or not isinstance(cell.value, str)
                or cell.value.strip() not in ALIASES
            ):
                raise AppError("WORKBOOK_UNSUPPORTED", "存在未知表头，请按模板提供字段。", 422)
            mapping.append(ALIASES[cell.value.strip()])
        actual = [field for field in mapping if field]
        if not REQUIRED.issubset(actual) or any(n > 1 for n in Counter(actual).values()):
            raise AppError("WORKBOOK_UNSUPPORTED", "必填列缺失或同一字段出现多次。", 422)
        rows = []
        for index, cells in enumerate(sheet.iter_rows(min_row=2, max_row=10001), 2):
            if index > 10000 and any(cell.value is not None for cell in cells):
                raise AppError("UPLOAD_TOO_LARGE", "工作表超过 10000 个物理行。", 413)
            if any(cell.data_type == "f" for cell in cells):
                raise AppError(
                    "WORKBOOK_UNSUPPORTED", "工作表含公式，请先将计算结果转换为普通数值。", 422
                )
            if not any(cell.value is not None for cell in cells):
                continue
            if len(rows) >= 2000:
                raise AppError("UPLOAD_TOO_LARGE", "每批最多 2000 条非空数据行。", 413)
            if any(cell.value is not None for cell in cells[len(mapping) :]):
                raise AppError("WORKBOOK_UNSUPPORTED", "有数据列缺少表头。", 422)
            values = {
                field: cells[i].value if i < len(cells) else None
                for i, field in enumerate(mapping)
                if field
            }
            if "company_id" not in mapping:
                values["company_id"] = company_id
            # Scope errors are never skippable, including on otherwise invalid rows.
            if values.get("company_id") != company_id:
                raise AppError(
                    "IMPORT_VALIDATION_FAILED",
                    "文件中有其他企业或未填写的企业编号，整批未处理。",
                    422,
                )
            try:
                normalized = _normalize(values, company_id)
                rows.append(
                    {
                        "row_number": index,
                        "normalized_values": normalized,
                        "classification": "new",
                        "reason_code": None,
                        "reason_message": None,
                    }
                )
            except (ValueError, TypeError, OverflowError) as exc:
                rows.append(
                    {
                        "row_number": index,
                        "normalized_values": {},
                        "classification": "invalid",
                        "reason_code": "INVALID_ROW",
                        "reason_message": str(exc),
                    }
                )
        if not rows:
            raise AppError("IMPORT_VALIDATION_FAILED", "发票工作表没有数据行。", 422)
        return rows
    except AppError:
        raise
    except Exception as exc:
        raise AppError(
            "WORKBOOK_UNSUPPORTED", "工作簿无法解析，请检查模板与文件是否损坏。", 422
        ) from exc
    finally:
        if workbook:
            workbook.close()


def file_hash(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()
